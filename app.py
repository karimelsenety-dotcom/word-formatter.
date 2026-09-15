import io
import re
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import streamlit as st

# الألوان المعتمدة
BLUE_COLOR = RGBColor(31, 73, 125)  # #1F497D
BLACK_COLOR = RGBColor(0, 0, 0)


# 1. ضبط اتجاه النص RTL داخل ملف Word
def set_rtl(paragraph_or_cell):
    pPr = paragraph_or_cell._element.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)


# 2. بناء ملف Word بالتنسيقات المحددة
def build_custom_docx(raw_text):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    lines = raw_text.strip().split("\n")

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # 1. العنوان الرئيسي: يبدأ بـ # (16 Bold - أزرق)
        if stripped.startswith("#"):
            clean_text = stripped.lstrip("#").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_rtl(p)

            run = p.add_run(clean_text)
            run.font.name = "Arial"
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = BLUE_COLOR
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(8)

        # 2. العنوان الفرعي: يبدأ بـ * (14 Bold - أزرق)
        elif stripped.startswith("*"):
            clean_text = stripped.lstrip("*").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_rtl(p)

            run = p.add_run(clean_text)
            run.font.name = "Arial"
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = BLUE_COLOR
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)

        # 3. النص العادي والنقاط: 12 عادي - أسود
        else:
            is_bullet = stripped.startswith("-") or stripped.startswith("•")
            clean_text = re.sub(r"^[-•]\s*", "", stripped)

            p = (
                doc.add_paragraph(style="List Bullet")
                if is_bullet
                else doc.add_paragraph()
            )
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_rtl(p)

            run = p.add_run(clean_text)
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.bold = False
            run.font.color.rgb = BLACK_COLOR
            p.paragraph_format.space_after = Pt(4)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# 3. إنشاء معاينة HTML بنفس القواعد
def generate_html_preview(raw_text):
    lines = raw_text.strip().split("\n")
    html_out = """
    <div style="
        background-color: #ffffff; 
        border: 1px solid #d1d5db; 
        border-radius: 8px; 
        padding: 30px; 
        font-family: Arial, sans-serif; 
        direction: rtl; 
        text-align: right; 
        box-shadow: 0px 4px 12px rgba(0,0,0,0.05);
        color: #000000;
        line-height: 1.8;
    ">
    """

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # 1. العنوان الرئيسي: يبدأ بـ #
        if stripped.startswith("#"):
            clean_text = stripped.lstrip("#").strip()
            html_out += f'<h1 style="color: #1F497D; font-size: 21px; font-weight: bold; margin-top: 16px; margin-bottom: 12px; border-bottom: 2px solid #1F497D; padding-bottom: 6px;">{clean_text}</h1>'

        # 2. العنوان الفرعي: يبدأ بـ *
        elif stripped.startswith("*"):
            clean_text = stripped.lstrip("*").strip()
            html_out += f'<h2 style="color: #1F497D; font-size: 18px; font-weight: bold; margin-top: 14px; margin-bottom: 6px; text-align: right;">{clean_text}</h2>'

        # 3. النص العادي
        else:
            is_bullet = stripped.startswith("-") or stripped.startswith("•")
            clean_text = re.sub(r"^[-•]\s*", "", stripped)
            prefix = "• " if is_bullet else ""
            margin = "margin-right: 15px;" if is_bullet else ""
            html_out += f'<p style="font-size: 16px; color: #000000; margin-bottom: 6px; font-weight: normal; {margin}">{prefix}{clean_text}</p>'

    html_out += "</div>"
    return html_out


# --- الواجهة ---
st.set_page_config(
    page_title="مُنسّق الملخصات بالمقاسات الخاصة", layout="wide"
)
st.title("📚 مُنسّق النصوص مع المعاينة المباشرة")

col1, col2 = st.columns([1, 1])

with col1:
    st.subheader("📝 إدخال النص")
    user_text = st.text_area(
        "ضع النص هنا:",
        height=450,
        placeholder="# العنوان الرئيسي\n* عنوان فرعي\nهذا نص عادي أو توضيحي...",
    )

with col2:
    st.subheader("👁️ معاينة الشاشة والتنسيق")
    if user_text.strip():
        preview_html = generate_html_preview(user_text)
        st.markdown(preview_html, unsafe_allow_html=True)
    else:
        st.info("أدخل النص في الجهة اليسرى لرؤية التنسيق والمعاينة مباشرة.")

st.divider()

if user_text.strip():
    file_data = build_custom_docx(user_text)
    st.download_button(
        label="📥 تحميل ملف Word المنسق (.docx)",
        data=file_data,
        file_name="ملخص_منسق.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary",
    )
